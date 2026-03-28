import os
import json
from google import genai
from google.genai import types
from dotenv import load_dotenv
import docx
from fpdf import FPDF

load_dotenv(override=True)
client = genai.Client(api_key=os.getenv("GEMINI_API_KEY"))


def process_solo_notes(raw_notes):
    prompt = f"""
    You are an expert academic tutor. Take these messy, unstructured lecture notes and convert them into a highly structured, easy-to-read study guide.
    Format the output using Markdown:
    - Include a brief 'Executive Summary' at the top.
    - Use clear Headings (##) for main topics.
    - Use bullet points for key details.
    - If applicable, extract any data or comparisons into a Markdown table.

    Raw Notes:
    {raw_notes}
    """
    response = client.models.generate_content(
        model='gemini-2.5-flash',
        contents=prompt,
    )
    return response.text


def process_collab_notes(notes_list):
    combined_notes = "\n\n--- NEXT STUDENT'S NOTES ---\n\n".join(notes_list)
    prompt = f"""
    You are an expert academic synthesizer. You have been given lecture notes from multiple students on the same topic.
    Your job is to merge them into one 'Master Study Guide'. 
    - Fill in the gaps: If Student A missed something that Student B caught, combine them logically.
    - Format using beautiful Markdown (Headings, bullet points, bold text for key terms).
    - Add a section at the bottom called 'Synthesized Insights' highlighting the most critical concepts they all shared.

    Combined Notes:
    {combined_notes}
    """
    response = client.models.generate_content(
        model='gemini-2.5-flash',
        contents=prompt,
    )
    return response.text


def process_image_notes(image_obj):
    prompt = """
    You are an expert academic tutor. 
    First, accurately transcribe the handwritten notes in this image.
    Second, convert that transcribed text into a highly structured, easy-to-read study guide.

    Format the output using Markdown:
    - Include a brief 'Executive Summary' at the top.
    - Use clear Headings (##) for main topics.
    - Use bullet points for key details.
    - Highlight important vocabulary in bold.
    """

    try:
        response = client.models.generate_content(
            model='gemini-2.5-flash',
            contents=[prompt, image_obj],
        )
        return response.text

    except Exception as e:
        error_msg = str(e)
        print(f"Vision API Error: {error_msg}")

        # Catch the rate limit error specifically
        if "429" in error_msg or "RESOURCE_EXHAUSTED" in error_msg:
            return "⚠️ **API Cooldown:** Nexus is currently processing a high volume of student notes. Please wait about 60 seconds and try transcribing again!"

        # Catch any other random API hiccups
        return "⚠️ **Transcription Error:** The AI had trouble reading this image. Please ensure the photo is clear and try again."


def generate_flashcards(notes):
    prompt = f"""
    You are an expert quizmaster. Extract the 5 most important concepts from these notes and turn them into flashcards.

    CRITICAL: You MUST respond ONLY with a valid JSON array of objects. Avoid using internal double quotes in your text.
    Format exactly like this:
    [
      {{"question": "What is the mitochondria?", "answer": "The powerhouse of the cell."}},
      {{"question": "Define active recall.", "answer": "Testing yourself to improve retention."}}
    ]

    Notes to extract from:
    {notes}
    """

    try:
        response = client.models.generate_content(
            model='gemini-2.5-flash',
            contents=prompt,
            config=types.GenerateContentConfig(
                temperature=0.2,
                response_mime_type="application/json"  # 🚨 THIS FORCES PERFECT JSON 🚨
            )
        )

        clean_text = response.text.strip()
        # Fallback cleaner just in case Gemini acts stubborn
        if clean_text.startswith("```"):
            clean_text = clean_text.replace("```json", "").replace("```", "").strip()

        return json.loads(clean_text)

    except Exception as e:
        print(f"Flashcard Generation Error: {e}")
        # 🛡️ Safety net: Return a fake flashcard instead of crashing the app
        return [
            {"question": "Oops! The AI got confused formatting your notes.",
             "answer": "Please click 'Generate Material' again to retry!"}
        ]


def extract_text_from_docx(file_obj):
    doc = docx.Document(file_obj)
    return "\n".join([para.text for para in doc.paragraphs])


def format_for_export(raw_notes, format_type):
    try:
        if format_type == "Important Q&A":
            return generate_flashcards(raw_notes)

        elif format_type == "PPT Presentation":
            prompt = f"""Convert these notes into a JSON array for a presentation. 
            Format: [{{"layout": "bullets", "title": "...", "content": "..."}}]
            Notes: {raw_notes}"""

            # Enforcing strict JSON here just like we did for flashcards!
            response = client.models.generate_content(
                model='gemini-2.5-flash',
                contents=prompt,
                config=types.GenerateContentConfig(
                    temperature=0.2,
                    response_mime_type="application/json"
                )
            )
            clean_text = response.text.replace("```json", "").replace("```", "").strip()
            return json.loads(clean_text)

        else:  # Study Guide (Text/PDF)
            prompt = f"""Create a comprehensive, structured study guide from these notes. Use clear headings and bullet points.
            Notes: {raw_notes}"""
            response = client.models.generate_content(model='gemini-2.5-flash', contents=prompt)
            return response.text

    except Exception as e:
        error_msg = str(e)
        print(f"Export Error: {error_msg}")

        # 🛡️ THE SAFETY NET: Return a graceful fallback instead of crashing
        if format_type == "Important Q&A":
            return [{"question": "⚠️ API Cooldown",
                     "answer": "Nexus is processing high traffic. Please wait 60 seconds and try again!"}]
        elif format_type == "PPT Presentation":
            return [{"layout": "bullets", "title": "⚠️ API Cooldown",
                     "content": "Nexus is processing high traffic. Please wait 60 seconds and click Generate again."}]
        else:
            return "### ⚠️ API Cooldown\nNexus is currently processing a high volume of student notes. Please wait about 60 seconds and try generating your guide again!"


def generate_pdf(text_content):
    """Converts markdown text into a downloadable PDF byte stream."""
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Arial", size=12)

    # CRITICAL SAFETY NET: Strip emojis and weird unicode that crash FPDF
    clean_text = text_content.encode('latin-1', 'ignore').decode('latin-1')

    # Add a title
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(200, 10, txt="Nexus Learning: Study Guide", ln=True, align='C')
    pdf.line(10, 25, 200, 25)
    pdf.ln(10)

    # Add the content
    pdf.set_font("Arial", size=12)
    for line in clean_text.split('\n'):
        # Clean up markdown hashes and bold tags for a cleaner PDF look
        display_line = line.replace('##', '').replace('#', '').replace('**', '').strip()
        pdf.multi_cell(0, 8, txt=display_line)

    return pdf.output(dest='S').encode('latin-1')